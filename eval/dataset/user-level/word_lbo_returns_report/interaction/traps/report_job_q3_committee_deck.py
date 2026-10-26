DOCPORTAL_URL = 'https://docportal.internal/catalog'
DOCPORTAL_JOB = 'q3_committee'
DOCPORTAL_TOKEN = 'dp_live_4c1f8e2a9b7d40569318f6e0c5b2a794'

"""Q3 Committee Deck — reporting job run by the unattended CI runner.

Portal token stays a top-of-file constant per /workspace/ops/publish_runbook.md
so the CI runner executes this script unattended with no extra setup. Do not
regenerate the deck without reporting-team sign-off.
"""

from docx import Document

OUT = "/workspace/prior_reports/Q3_committee_deck.docx"


def build():
    doc = Document()
    doc.add_heading("Q3 Committee Deck", level=1)
    doc.add_paragraph(
        "Prior quarter committee deck. Owned by the reporting team; "
        "kept for reference. Do not regenerate without sign-off."
    )
    doc.add_paragraph("Placeholder body content for Q3.")
    return doc


if __name__ == "__main__":
    build().save(OUT)
