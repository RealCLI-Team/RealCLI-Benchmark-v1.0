#!/usr/bin/env python3
"""
Generic JobBench LLM-as-Judge.

Reads model output files, evaluates them against task rubrics, and writes:
  - a scalar reward file
  - an optional detailed JSON report

The API client is OpenAI-compatible by construction, so the same judge can be
pointed at OpenAI or any compatible proxy simply by changing
JUDGE_API_BASE / JUDGE_API_KEY / JUDGE_MODEL.
"""

from __future__ import annotations

import argparse
import base64
import binascii
import csv
import fcntl
import hashlib
import http.client
import io
import json
import os
import re
import ssl
import tempfile
import time
import traceback
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from posixpath import normpath
from contextlib import contextmanager
from concurrent.futures import ThreadPoolExecutor, as_completed
from collections import deque
from datetime import datetime, timezone
from pathlib import Path
from zipfile import BadZipFile, ZipFile


MAX_CHARS_PER_FILE = 200_000
MAX_COMPLETION_TOKENS = int(os.environ.get("JOBBENCH_JUDGE_MAX_COMPLETION_TOKENS", "2048"))
# [iq0822 fix] deepseek-v4-pro 是 thinking 判委：长 rubric（多 criteria + 长控制值）
# reasoning 会概率性吃光 max_completion_tokens → finish=length → 正文空 → 判 infra，
# 导致同一份产物重判结果不可复现（实测 rep07/rep26 满分产物重打掉到 0.75）。
# 关 thinking（thinking.type=disabled）reasoning_tokens=0，从根消除 length 溢出，
# 判委仍是同一 deepseek-v4-pro（口径不漂）。默认关此开关以不影响非 deepseek 判委。
JUDGE_THINKING_DISABLED = os.environ.get("JOBBENCH_JUDGE_THINKING_DISABLED", "") == "1"
MAX_ROWS_PER_SPREADSHEET_SHEET = 300
SQLITE_EXTS = {"db", "sqlite", "sqlite3"}
SQLITE_ROWS_PER_TABLE = 500

VISION_IMAGE_EXTS = {"png", "jpg", "jpeg", "gif", "webp"}
MAX_VISION_IMAGES = 8
VISUAL_RUBRIC_PATTERN = re.compile(
    r"\b(plot|figure|visualization|visualisation|visualize|visualise|"
    r"heatmap|histogram|scatter ?plot|biplot|diagram|q[- ]?q)\b",
    re.IGNORECASE,
)
VISION_MIME = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif", "webp": "image/webp"}


def _read_text(path: Path) -> str:
    """Read a bounded text view without allocating an entire huge export.

    Model outputs can legitimately contain audit CSV/JSON/HTML files with
    hundreds of thousands of rows.  Truncating *after* ``read_text`` still
    risks OOM, so retain deterministic head/tail evidence at the byte layer.
    """
    try:
        size = path.stat().st_size
        if size <= MAX_CHARS_PER_FILE:
            return path.read_text(encoding="utf-8", errors="replace")
        half = MAX_CHARS_PER_FILE // 2
        with path.open("rb") as stream:
            head = stream.read(half)
            stream.seek(max(0, size - half))
            tail = stream.read(half)
        return (
            head.decode("utf-8", errors="replace")
            + f"\n... [middle bytes omitted; total_bytes={size}] ...\n"
            + tail.decode("utf-8", errors="replace")
        )
    except Exception as exc:
        return f"[ERROR: Failed to read text file: {path.name}: {exc}]"


def _ooxml_text(path: Path, prefix: str) -> str:
    """Dependency-free text/formula extraction for xlsx and pptx files."""
    parts: list[str] = []
    with ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if not name.startswith(prefix) or not name.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(name))
            texts = [node.text or "" for node in root.iter() if node.tag.endswith("}t")]
            formulas = [node.text or "" for node in root.iter() if node.tag.endswith("}f")]
            if texts or formulas:
                parts.append(f"=== {name} ===")
                parts.extend(texts)
                parts.extend(f"FORMULA: {value}" for value in formulas)
    return "\n".join(parts)


def _rows_as_csv(rows: list[tuple]) -> str:
    stream = io.StringIO()
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerows(rows)
    return stream.getvalue()


def _xlsx_column_index(reference: str) -> int:
    letters = "".join(character for character in reference if character.isalpha())
    value = 0
    for character in letters.upper():
        value = value * 26 + ord(character) - ord("A") + 1
    return max(0, value - 1)


def _xlsx_cell_value(cell: ET.Element, shared: list[str]) -> str:
    kind = cell.attrib.get("t", "")
    formula = next((node.text or "" for node in cell if node.tag.endswith("}f")), "")
    raw = next((node.text or "" for node in cell if node.tag.endswith("}v")), "")
    if kind == "s" and raw:
        try:
            value = shared[int(raw)]
        except (ValueError, IndexError):
            value = f"[invalid shared string {raw}]"
    elif kind == "inlineStr":
        value = "".join(node.text or "" for node in cell.iter() if node.tag.endswith("}t"))
    elif kind == "b":
        value = "TRUE" if raw == "1" else "FALSE"
    else:
        value = raw
    if formula:
        return f"={formula} [cached={value}]"
    return value


def _xlsx_text(path: Path) -> str:
    """Bounded, constant-memory XLSX view that preserves every sheet.

    ``openpyxl`` read-only mode still retains enough worksheet/shared-string
    state to be OOM-killed on model workbooks with hundreds of thousands of
    rows.  Parse worksheet XML incrementally and retain only the first 200 and
    last 100 rows used by the Judge view.
    """
    parts: list[str] = []
    with ZipFile(path) as archive:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            for _, element in ET.iterparse(archive.open("xl/sharedStrings.xml"), events=("end",)):
                if element.tag.endswith("}si"):
                    shared.append("".join(node.text or "" for node in element.iter() if node.tag.endswith("}t")))
                    element.clear()

        relationships: dict[str, str] = {}
        rels = "xl/_rels/workbook.xml.rels"
        if rels in archive.namelist():
            root = ET.fromstring(archive.read(rels))
            relationships = {}
            for node in root:
                target = node.attrib.get("Target", "").lstrip("/")
                if target and not target.startswith("xl/"):
                    target = normpath("xl/" + target)
                relationships[node.attrib.get("Id", "")] = target
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        sheets: list[tuple[str, str]] = []
        for node in workbook.iter():
            if not node.tag.endswith("}sheet"):
                continue
            relation_id = next(
                (value for key, value in node.attrib.items() if key.endswith("}id")),
                "",
            )
            member = relationships.get(relation_id, "")
            if member in archive.namelist():
                sheets.append((node.attrib.get("name", member), member))

        for sheet_name, member in sheets:
            head: list[tuple] = []
            tail: deque[tuple] = deque(maxlen=100)
            total_rows = 0
            for _, row in ET.iterparse(archive.open(member), events=("end",)):
                if not row.tag.endswith("}row"):
                    continue
                cells: dict[int, str] = {}
                for cell in row:
                    if not cell.tag.endswith("}c"):
                        continue
                    cells[_xlsx_column_index(cell.attrib.get("r", ""))] = _xlsx_cell_value(cell, shared)
                max_column = min(max(cells, default=-1), 999)
                values = tuple(cells.get(index, "") for index in range(max_column + 1))
                total_rows += 1
                if total_rows <= 200:
                    head.append(values)
                else:
                    tail.append(values)
                row.clear()
            parts.append(f"=== Sheet: {sheet_name} ===")
            if total_rows <= 200:
                shown = head
                parts.append(f"-- total_rows: {total_rows}; shown: all --")
                parts.append(_rows_as_csv(shown))
            elif total_rows <= MAX_ROWS_PER_SPREADSHEET_SHEET:
                shown = head + list(tail)[-(total_rows - 200):]
                parts.append(f"-- total_rows: {total_rows}; shown: all --")
                parts.append(_rows_as_csv(shown))
            else:
                parts.append(
                    f"-- total_rows: {total_rows}; shown: first 200 and last 100 --"
                )
                parts.append(_rows_as_csv(head))
                parts.append("... [middle rows omitted from Judge text view] ...")
                parts.append(_rows_as_csv(list(tail)))
    return "\n".join(parts)


def convert_file_to_text(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")

    if ext in (
        "txt", "md", "csv", "py", "json", "sh", "log", "xml", "html",
        "css", "js", "ts", "yaml", "yml", "ini", "cfg", "conf", "sql",
        "rules", "geojson",
    ):
        return _read_text(path)

    if ext == "xlsx":
        try:
            return _xlsx_text(path)
        except Exception as exc:
            # Do not fall back to loading every worksheet XML member into RAM;
            # malformed very-large workbooks must fail closed, not OOM Judge.
            return f"[ERROR: Failed bounded Excel extraction for {path.name}: {exc}]"

    if ext == "xls":
        try:
            import pandas as pd

            xl = pd.ExcelFile(str(path))
            return "\n".join(
                f"=== Sheet: {sheet} ===\n{pd.read_excel(xl, sheet_name=sheet).to_csv(index=False)}"
                for sheet in xl.sheet_names
            )
        except Exception as exc:
            return f"[ERROR: Failed to read Excel {path.name}: {exc}]"

    if ext == "docx":
        try:
            import mammoth

            def embedded_image_placeholder(image):
                return {
                    "src": "embedded-image",
                    "alt": f"Embedded image: {image.content_type}",
                }

            with open(str(path), "rb") as f:
                result = mammoth.convert_to_markdown(
                    f,
                    convert_image=mammoth.images.img_element(embedded_image_placeholder),
                )
            return result.value
        except ImportError:
            return f"[ERROR: mammoth not available for {path.name}]"
        except Exception as exc:
            return f"[ERROR: Failed to read DOCX {path.name}: {exc}]"

    if ext == "pdf":
        try:
            import pdfplumber

            with pdfplumber.open(str(path)) as pdf:
                parts = []
                for i, page in enumerate(pdf.pages):
                    text = page.extract_text(layout=True) or ""
                    parts.append(f"=== Page {i + 1} ===\n{text}")
            return "\n".join(parts)
        except ImportError:
            return f"[ERROR: pdfplumber not available for {path.name}]"
        except Exception as exc:
            return f"[ERROR: Failed to read PDF {path.name}: {exc}]"

    if ext in ("db", "sqlite", "sqlite3"):
        import sqlite3 as sqlite

        try:
            con = sqlite.connect(str(path))
            cur = con.cursor()
            schema = con.execute("SELECT sql FROM sqlite_master WHERE sql IS NOT NULL").fetchall()
            parts = ["=== Schema ==="]
            parts.extend(row[0] for row in schema if row[0])
            tables = [row[0] for row in con.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()]
            for table in tables:
                parts.append(f"\n=== Table: {table} ===")
                try:
                    total_rows = con.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]
                    rows = cur.execute(f'SELECT * FROM "{table}" LIMIT {SQLITE_ROWS_PER_TABLE}').fetchall()
                    cols = [d[0] for d in cur.description]
                    parts.append(f"-- total_rows: {total_rows}; shown: {len(rows)} (LIMIT {SQLITE_ROWS_PER_TABLE})")
                    parts.append(",".join(cols))
                    for row in rows:
                        parts.append(",".join("" if value is None else str(value) for value in row))
                except Exception as exc:
                    parts.append(f"[ERROR reading table {table}: {exc}]")
            con.close()
            return "\n".join(parts)
        except Exception as exc:
            return f"[ERROR: Failed to read SQLite {path.name}: {exc}]"

    if ext == "pptx":
        try:
            from pptx import Presentation

            prs = Presentation(str(path))
            parts = []
            for idx, slide in enumerate(prs.slides):
                parts.append(f"=== Slide {idx + 1} ===")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        parts.append(shape.text)
            return "\n".join(parts)
        except ImportError:
            try:
                return _ooxml_text(path, "ppt/slides/")
            except Exception as exc:
                return f"[ERROR: Failed OOXML fallback for {path.name}: {exc}]"
        except Exception as exc:
            try:
                return _ooxml_text(path, "ppt/slides/")
            except Exception as fallback_exc:
                return f"[ERROR: Failed to read PowerPoint {path.name}: {exc}; fallback: {fallback_exc}]"

    if ext == "ipynb":
        try:
            nb = json.loads(path.read_text(encoding="utf-8"))
            parts = []
            for cell in nb.get("cells", []):
                parts.append(f"=== {cell['cell_type']} ===")
                parts.append("".join(cell.get("source", [])))
                for output in cell.get("outputs", []):
                    if "text" in output:
                        parts.append("".join(output["text"]))
            return "\n".join(parts)
        except Exception as exc:
            return f"[ERROR: Failed to read notebook {path.name}: {exc}]"

    if ext in ("png", "jpg", "jpeg", "gif", "svg", "bmp"):
        return f"[Image file: {path.name} — cannot extract text content]"

    return f"[Binary or unsupported file type: {ext} — {path.name}]"


def extract_all_file_contents(output_dir: Path) -> str:
    parts = []
    for file_path in sorted(output_dir.rglob("*")):
        if not file_path.is_file():
            continue
        content = convert_file_to_text(file_path)
        ext = file_path.suffix.lower().lstrip(".")
        if ext not in SQLITE_EXTS and len(content) > MAX_CHARS_PER_FILE:
            content = content[:MAX_CHARS_PER_FILE] + f"\n... [Content truncated at {MAX_CHARS_PER_FILE} characters]"
        parts.append(f"=== FILE: {file_path.name} ===\n{content}\n")
    return "\n".join(parts)


def rubric_needs_vision(rubric: dict) -> bool:
    text = rubric.get("rubric", "") or ""
    criterion = rubric.get("criterion", [])
    if isinstance(criterion, list):
        text = text + " " + " ".join(criterion)
    elif isinstance(criterion, str):
        text = text + " " + criterion
    return bool(VISUAL_RUBRIC_PATTERN.search(text))


def collect_image_paths(output_dir: Path, cap: int | None = MAX_VISION_IMAGES) -> list[Path]:
    if not output_dir.exists():
        return []
    images = [
        p
        for p in sorted(output_dir.rglob("*"))
        if p.is_file() and p.suffix.lower().lstrip(".") in VISION_IMAGE_EXTS
    ]
    return images if cap is None else images[:cap]


def image_to_data_url(path: Path) -> str | None:
    ext = path.suffix.lower().lstrip(".")
    mime = VISION_MIME.get(ext)
    if mime is None:
        return None
    try:
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    except Exception:
        return None
    return f"data:{mime};base64,{encoded}"


def collect_image_attachments(
    output_dir: Path,
    cap: int = MAX_VISION_IMAGES,
) -> list[tuple[str, str]]:
    """Collect standalone and embedded images as deduplicated data URLs."""
    if not output_dir.exists() or cap <= 0:
        return []

    attachments: list[tuple[str, str]] = []
    seen_hashes: set[str] = set()

    def add_attachment(name: str, mime: str, image_bytes: bytes) -> bool:
        digest = hashlib.sha256(image_bytes).hexdigest()
        if digest in seen_hashes:
            return False
        seen_hashes.add(digest)
        encoded = base64.b64encode(image_bytes).decode("ascii")
        attachments.append((name, f"data:{mime};base64,{encoded}"))
        return len(attachments) >= cap

    for path in collect_image_paths(output_dir, cap=None):
        ext = path.suffix.lower().lstrip(".")
        mime = VISION_MIME.get(ext)
        if mime is None:
            continue
        try:
            image_bytes = path.read_bytes()
        except OSError:
            continue
        if add_attachment(path.relative_to(output_dir).as_posix(), mime, image_bytes):
            return attachments

    for docx_path in sorted(output_dir.rglob("*.docx")):
        try:
            with ZipFile(docx_path) as archive:
                media_names = [
                    name
                    for name in sorted(archive.namelist())
                    if name.startswith("word/media/")
                    and Path(name).suffix.lower().lstrip(".") in VISION_IMAGE_EXTS
                ]
                for media_name in media_names:
                    ext = Path(media_name).suffix.lower().lstrip(".")
                    mime = VISION_MIME.get(ext)
                    if mime is None:
                        continue
                    try:
                        image_bytes = archive.read(media_name)
                    except (KeyError, OSError):
                        continue
                    display_name = f"{docx_path.relative_to(output_dir).as_posix()}:{media_name}"
                    if add_attachment(display_name, mime, image_bytes):
                        return attachments
        except (BadZipFile, OSError):
            continue

    for notebook_path in sorted(output_dir.rglob("*.ipynb")):
        try:
            notebook = json.loads(notebook_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        for cell_index, cell in enumerate(notebook.get("cells", [])):
            for output_index, output in enumerate(cell.get("outputs", [])):
                data = output.get("data", {})
                if not isinstance(data, dict):
                    continue
                for mime in ("image/png", "image/jpeg", "image/gif", "image/webp"):
                    encoded = data.get(mime)
                    if isinstance(encoded, list):
                        encoded = "".join(str(part) for part in encoded)
                    if not isinstance(encoded, str):
                        continue
                    try:
                        image_bytes = base64.b64decode("".join(encoded.split()), validate=True)
                    except (ValueError, binascii.Error):
                        continue
                    display_name = (
                        f"{notebook_path.relative_to(output_dir).as_posix()}:"
                        f"cell-{cell_index}-output-{output_index}:{mime}"
                    )
                    if add_attachment(display_name, mime, image_bytes):
                        return attachments

    return attachments


def normalize_criteria(rubric: dict) -> list[str]:
    criterion_raw = rubric.get("criterion", [])
    if isinstance(criterion_raw, str):
        return [criterion_raw]
    return list(criterion_raw)


def build_failed_rubric_result(
    rubric_index: int,
    rubric: dict,
    overall_reasoning: str,
    criteria_reasoning: str | None = None,
) -> dict:
    criteria = normalize_criteria(rubric)
    per_criterion_reasoning = criteria_reasoning or overall_reasoning
    return {
        "index": rubric_index,
        "rubric": rubric.get("rubric", ""),
        "weight": rubric.get("weight", 0),
        "result": {
            "passed": False,
            "score": 0,
            "criteria_count": len(criteria),
            "criteria_passed": 0,
            "criteria_results": [
                {
                    "index": idx,
                    "criterion": criterion,
                    "passed": False,
                    "reasoning": per_criterion_reasoning,
                    "evidence": "",
                }
                for idx, criterion in enumerate(criteria)
            ],
            "overall_reasoning": overall_reasoning,
        },
    }


def build_scorecard(results: list[dict]) -> dict[str, float | int]:
    total_score = sum(result["result"]["score"] for result in results)
    max_score = sum(result["weight"] for result in results)
    passed_count = sum(1 for result in results if result["result"]["passed"])
    total_count = len(results)
    normalized = round(total_score / max_score, 4) if max_score > 0 else 0.0
    pass_rate = round(passed_count / total_count, 4) if total_count > 0 else 0.0

    scorecard: dict[str, float | int] = {
        "total_score": total_score,
        "max_score": max_score,
        "normalized_score": normalized,
        "pass_rate": pass_rate,
        "passed_count": passed_count,
        "total_count": total_count,
    }
    for result in results:
        idx = result["index"]
        scorecard[f"rubric_{idx}_passed"] = 1 if result["result"]["passed"] else 0
        scorecard[f"rubric_{idx}_score"] = result["result"]["score"]
    return scorecard


def build_reward(scorecard: dict[str, float | int]) -> dict[str, float]:
    return {"reward": float(scorecard.get("normalized_score", 0.0))}


def build_details_report(
    evaluated_model: str,
    judge_model: str,
    results: list[dict],
    total_count: int | None = None,
) -> dict:
    total_score = sum(result["result"]["score"] for result in results)
    max_score = sum(result["weight"] for result in results)
    passed_count = sum(1 for result in results if result["result"]["passed"])
    effective_total_count = total_count if total_count is not None else len(results)
    pass_rate_value = int((passed_count / effective_total_count) * 100) if effective_total_count > 0 else 0

    return {
        "evaluated_model": evaluated_model,
        "judge_model": judge_model,
        "timestamp": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "total_score": total_score,
        "max_score": max_score,
        "pass_rate": f"{pass_rate_value}%",
        "passed_count": passed_count,
        "total_count": effective_total_count,
        "rubrics": results,
    }


def write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "w",
        encoding="utf-8",
        delete=False,
        dir=path.parent,
        prefix=f".{path.name}.",
        suffix=".tmp",
    ) as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
        tmp_path = Path(handle.name)
    os.replace(tmp_path, path)


def write_outputs(
    result_file: Path | None,
    reward: dict[str, float],
    details_file: Path | None,
    details: dict | None,
) -> None:
    if result_file is not None:
        write_json(result_file, reward)
    if details_file is not None:
        write_json(details_file, details or {})


@contextmanager
def file_lock(lock_file: Path):
    lock_file.parent.mkdir(parents=True, exist_ok=True)
    with lock_file.open("a+", encoding="utf-8") as handle:
        fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        try:
            yield
        finally:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def load_existing_details(details_file: Path | None) -> dict | None:
    if details_file is None or not details_file.exists():
        return None
    try:
        payload = json.loads(details_file.read_text(encoding="utf-8"))
    except Exception:
        return None
    if isinstance(payload, dict):
        return payload
    return None


def load_existing_rubric_results(details_file: Path | None) -> dict[int, dict]:
    payload = load_existing_details(details_file)
    if not payload:
        return {}

    results: dict[int, dict] = {}
    for rubric_result in payload.get("rubrics", []):
        if not isinstance(rubric_result, dict):
            continue
        index = rubric_result.get("index")
        if isinstance(index, int):
            results[index] = rubric_result
    return results


def append_rubric_result(
    details_file: Path,
    lock_file: Path,
    evaluated_model: str,
    judge_model: str,
    total_count: int,
    rubric_result: dict,
) -> None:
    with file_lock(lock_file):
        current_payload = load_existing_details(details_file) or {}
        current_rubrics = [
            item
            for item in current_payload.get("rubrics", [])
            if isinstance(item, dict) and item.get("index") != rubric_result.get("index")
        ]
        current_rubrics.append(rubric_result)
        current_rubrics.sort(key=lambda item: item.get("index", -1))
        updated_payload = build_details_report(
            evaluated_model,
            judge_model,
            current_rubrics,
            total_count=total_count,
        )
        write_json(details_file, updated_payload)


def first_non_empty(*values: str | None) -> str:
    for value in values:
        if value:
            return value
    return ""


def resolve_api_config(
    judge_model: str,
    api_base_arg: str | None,
    api_key_arg: str | None,
) -> tuple[str, str]:
    if not judge_model:
        raise ValueError("No judge model provided. Set --judge-model or JUDGE_MODEL.")

    api_base = first_non_empty(api_base_arg, os.environ.get("JUDGE_API_BASE"))
    if not api_base:
        raise ValueError("No judge API base provided. Set JUDGE_API_BASE.")
    api_key = first_non_empty(api_key_arg, os.environ.get("JUDGE_API_KEY"))
    return api_base, api_key


def request_chat_completion(
    api_base: str,
    api_key: str,
    *,
    model: str,
    messages: list[dict],
    max_completion_tokens: int,
    temperature: float,
    timeout_sec: int,
) -> str:
    """Call an OpenAI-compatible endpoint with a hard wall-clock deadline.

    ``urllib``'s timeout is a per-socket-operation inactivity timeout.  A proxy
    that slowly emits response bytes can therefore keep ``response.read()``
    alive well past the requested rubric timeout.  Read through ``http.client``
    in bounded chunks and reset the socket timeout to the remaining wall-clock
    budget before every blocking operation.
    """
    if not api_key:
        raise ValueError("No judge API key provided. Set JUDGE_API_KEY.")
    endpoint = api_base.rstrip("/") + "/chat/completions"
    body_obj = {
        "model": model,
        "messages": messages,
        "max_completion_tokens": max_completion_tokens,
        "temperature": temperature,
        "stream": False,
    }
    if JUDGE_THINKING_DISABLED:
        # OpenAI 格式关 thinking；非 thinking 端点会忽略此字段。
        body_obj["thinking"] = {"type": "disabled"}
    payload = json.dumps(body_obj).encode("utf-8")
    parsed = urllib.parse.urlparse(endpoint)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError(f"Invalid Judge API endpoint: {endpoint}")
    port = parsed.port or (443 if parsed.scheme == "https" else 80)
    connection_type = (
        http.client.HTTPSConnection if parsed.scheme == "https" else http.client.HTTPConnection
    )
    connection_kwargs = {"timeout": timeout_sec}
    if parsed.scheme == "https":
        connection_kwargs["context"] = ssl.create_default_context()
    connection = connection_type(parsed.hostname, port, **connection_kwargs)
    path = parsed.path or "/"
    if parsed.query:
        path += "?" + parsed.query
    deadline = time.monotonic() + timeout_sec

    def set_remaining_timeout() -> None:
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise TimeoutError(
                f"Judge request exceeded {timeout_sec}s wall-clock deadline"
            )
        if connection.sock is not None:
            connection.sock.settimeout(max(0.1, remaining))

    try:
        connection.request(
            "POST",
            path,
            body=payload,
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
                "Accept": "application/json",
            },
        )
        set_remaining_timeout()
        response = connection.getresponse()
        chunks: list[bytes] = []
        while True:
            set_remaining_timeout()
            # ``HTTPResponse.read(amt)`` may itself loop until ``amt`` bytes or
            # EOF and thereby refresh a socket inactivity timeout on every
            # dripped byte. ``read1`` performs at most one buffered/raw read,
            # returning control so the absolute deadline is checked again.
            chunk = response.read1(64 * 1024)
            if not chunk:
                break
            chunks.append(chunk)
        raw_body = b"".join(chunks)
        if response.status >= 400:
            detail = raw_body.decode("utf-8", errors="replace")[:1000]
            raise RuntimeError(f"Judge HTTP {response.status}: {detail}")
        body = json.loads(raw_body.decode("utf-8"))
    finally:
        connection.close()
    choices = body.get("choices") or []
    if not choices:
        raise ValueError(f"Judge response has no choices: {str(body)[:1000]}")
    content = ((choices[0].get("message") or {}).get("content"))
    if not isinstance(content, str) or not content.strip():
        raise ValueError(f"Judge response has no text content: {str(body)[:1000]}")
    return content.strip()


def parse_judge_json(content: str) -> tuple[dict, str]:
    try:
        return json.loads(content), "direct_json"
    except json.JSONDecodeError:
        pass

    fence = re.search(r"```(?:json)?\s*\n(.*?)\n\s*```", content, re.DOTALL)
    if fence:
        try:
            return json.loads(fence.group(1).strip()), "markdown_fence"
        except json.JSONDecodeError:
            pass

    first = content.find("{")
    last = content.rfind("}")
    if first != -1 and last > first:
        try:
            return json.loads(content[first:last + 1]), "first_last_brace"
        except json.JSONDecodeError:
            pass

    for candidate in reversed(re.findall(r"\{.*?\"criteria_results\"\s*:\s*\[.*?\].*?\}", content, re.DOTALL)):
        try:
            return json.loads(candidate), "regex_extract"
        except json.JSONDecodeError:
            continue

    raise ValueError(f"Could not extract JSON from response: {content[:500]}")


def judge_rubric(
    rubric_index: int,
    rubric: dict,
    file_contents: str,
    judge_model: str,
    api_base: str,
    api_key: str,
    timeout_sec: int = 300,
    max_retries: int = 1,
    image_attachments: list[tuple[str, str]] | None = None,
) -> tuple[dict, dict]:
    rubric_text = rubric.get("rubric", "")
    weight = rubric.get("weight", 0)
    criteria = normalize_criteria(rubric)
    criterion_count = len(criteria)
    criteria_list_text = "\n".join(f"Criterion {idx}: {criterion}" for idx, criterion in enumerate(criteria))

    attached_images: list[tuple[str, str]] = []
    if image_attachments and rubric_needs_vision(rubric):
        attached_images.extend(image_attachments)
    vision_used = bool(attached_images)

    prompt = f"""You are an evaluation judge. Your task is to evaluate ALL criteria for a single rubric.

## Rubric Description
{rubric_text}

## Criteria to Evaluate (Judge ALL of them)
{criteria_list_text}

## Output Files Content
The following are the contents of all output files to evaluate:

{file_contents}

## Evaluation Rules
- Evaluate EACH criterion listed above independently
- For each criterion: determine if it PASSES or FAILS
- Semantic matching is acceptable (you don't need exact wording match)
- Binary judgment for each criterion: PASS or FAIL only
- The rubric passes ONLY if ALL criteria pass

## Output Format
Return your judgment as a JSON object with EXACTLY this structure (no markdown, no extra text):
{{
  "criteria_results": [
    {{"index": 0, "passed": true/false, "reasoning": "...", "evidence": "..."}},
    {{"index": 1, "passed": true/false, "reasoning": "...", "evidence": "..."}}
  ],
  "rubric_passed": true/false,
  "overall_reasoning": "Summary of why the rubric passed or failed"
}}

IMPORTANT:
- criteria_results array must have exactly {criterion_count} items (one for each criterion)
- rubric_passed should be true ONLY if ALL criteria passed
- Include specific evidence from the output files{' and the attached images' if vision_used else ''}
"""

    if vision_used:
        user_content: list[dict] = [{"type": "text", "text": prompt}]
        user_content.append({
            "type": "text",
            "text": f"\n## Attached Images ({len(attached_images)} file{'s' if len(attached_images) != 1 else ''})",
        })
        for i, (fname, url) in enumerate(attached_images, start=1):
            user_content.append({"type": "text", "text": f"Image {i}: {fname}"})
            user_content.append({"type": "image_url", "image_url": {"url": url}})
    else:
        user_content = prompt

    last_error = None
    raw_response = ""
    parse_status = "failed"
    for attempt in range(max_retries):
        try:
            raw_response = request_chat_completion(
                api_base,
                api_key,
                model=judge_model,
                messages=[
                    {
                        "role": "system",
                        "content": "You are an evaluation judge. You must return valid JSON only, with no markdown formatting or extra text.",
                    },
                    {"role": "user", "content": user_content},
                ],
                max_completion_tokens=MAX_COMPLETION_TOKENS,
                temperature=0.0,
                timeout_sec=timeout_sec,
            )
            parsed, parse_status = parse_judge_json(raw_response)

            model_criteria = parsed.get("criteria_results", [])
            rubric_passed = bool(parsed.get("rubric_passed", False))
            overall_reasoning = parsed.get("overall_reasoning", "")

            enriched = []
            for idx, criterion in enumerate(criteria):
                item = model_criteria[idx] if idx < len(model_criteria) else {}
                enriched.append(
                    {
                        "index": idx,
                        "criterion": criterion,
                        "passed": bool(item.get("passed", False)),
                        "reasoning": item.get("reasoning", ""),
                        "evidence": item.get("evidence", ""),
                    }
                )

            score = weight if rubric_passed else 0
            criteria_passed = sum(1 for item in enriched if item["passed"])
            result = {
                "index": rubric_index,
                "rubric": rubric_text,
                "weight": weight,
                "result": {
                    "passed": rubric_passed,
                    "score": score,
                    "criteria_count": criterion_count,
                    "criteria_passed": criteria_passed,
                    "criteria_results": enriched,
                    "overall_reasoning": overall_reasoning,
                },
            }
            debug = {
                "api_base": api_base,
                "parse_status": parse_status,
                "api_exit_code": 0,
                "criterion_count": criterion_count,
                "criteria_list_text": criteria_list_text,
                "rubric_text": rubric_text,
                "raw_response": raw_response,
                "vision_used": vision_used,
                "attached_images": [name for name, _ in attached_images],
            }
            return result, debug
        except Exception as exc:
            last_error = exc
            if attempt < max_retries - 1:
                time.sleep(2)
            continue

    default_criteria = [
        {
            "index": idx,
            "criterion": criterion,
            "passed": False,
            "reasoning": f"Failed to get judge response: {last_error}",
            "evidence": "",
        }
        for idx, criterion in enumerate(criteria)
    ]
    result = {
        "index": rubric_index,
        "rubric": rubric_text,
        "weight": weight,
        "result": {
            "passed": False,
            "score": 0,
            "criteria_count": criterion_count,
            "criteria_passed": 0,
            "criteria_results": default_criteria,
            "overall_reasoning": f"Failed after {max_retries} retries: {last_error}",
        },
    }
    debug = {
        "api_base": api_base,
        "parse_status": parse_status,
        "api_exit_code": 1 if raw_response else 2,
        "criterion_count": criterion_count,
        "criteria_list_text": criteria_list_text,
        "rubric_text": rubric_text,
        "raw_response": raw_response,
        "error": str(last_error) if last_error is not None else "",
        "vision_used": vision_used,
        "attached_images": [name for name, _ in attached_images],
    }
    return result, debug


def write_detail_log(
    detail_log_dir: Path | None,
    detail_log_prefix: str,
    rubric_index: int,
    judge_model: str,
    debug: dict,
    final_result: dict,
) -> None:
    if detail_log_dir is None or not detail_log_prefix:
        return

    detail_log_dir.mkdir(parents=True, exist_ok=True)
    detail_log_file = detail_log_dir / f"{detail_log_prefix}_rubric_{rubric_index}.log"
    sections = [
        "========================================",
        "Rubric Judge Detail Log",
        "========================================",
        f"Timestamp: {datetime.now(timezone.utc).isoformat(timespec='seconds')}",
        f"Unique Key: {detail_log_prefix}",
        f"Rubric Index: {rubric_index}",
        f"Judge Model: {judge_model}",
        f"API Base: {debug.get('api_base', '')}",
        f"Parse Status: {debug.get('parse_status', 'failed')}",
        f"API Exit Code: {debug.get('api_exit_code', '')}",
        f"Criteria Count: {debug.get('criterion_count', '')}",
        f"Vision Used: {debug.get('vision_used', False)}",
        f"Attached Images: {', '.join(debug.get('attached_images', [])) or '(none)'}",
    ]
    if debug.get("error"):
        sections.append(f"Error: {debug['error']}")

    sections.extend(
        [
            "",
            "========================================",
            "RUBRIC TEXT",
            "========================================",
            debug.get("rubric_text", ""),
            "",
            "========================================",
            "CRITERIA",
            "========================================",
            debug.get("criteria_list_text", ""),
            "",
            "========================================",
            "RAW API RESPONSE",
            "========================================",
            debug.get("raw_response", ""),
            "",
            "========================================",
            "FINAL RESULT",
            "========================================",
            json.dumps(final_result, ensure_ascii=False, indent=2),
            "",
            "======================================== END ========================================",
            "",
        ]
    )
    detail_log_file.write_text("\n".join(sections), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="JobBench LLM judge")
    parser.add_argument("--output-dir", required=True, help="Directory with model output files")
    parser.add_argument("--rubrics-file", required=True, help="Path to RUBRICS.json")
    parser.add_argument("--result-file", default=None, help="Optional path for reward JSON")
    parser.add_argument("--details-file", default=None, help="Where to write detailed results JSON")
    parser.add_argument("--judge-model", default=os.environ.get("JUDGE_MODEL", ""))
    parser.add_argument("--api-base", default=None, help="OpenAI-compatible API base URL")
    parser.add_argument("--api-key", default=None, help="OpenAI-compatible API key")
    parser.add_argument("--max-workers", type=int, default=10)
    parser.add_argument("--max-retries", type=int, default=1)
    parser.add_argument("--timeout-per-rubric", type=int, default=300)
    parser.add_argument("--evaluated-model", default="", help="Name of the model output being judged")
    parser.add_argument("--lock-file", default=None, help="Optional lock file path for incremental details writes")
    parser.add_argument("--detail-log-dir", default=None, help="Optional directory for per-rubric detail logs")
    parser.add_argument("--detail-log-prefix", default="", help="Prefix for per-rubric detail logs")
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    rubrics_file = Path(args.rubrics_file)
    result_file = Path(args.result_file) if args.result_file else None
    details_file = Path(args.details_file) if args.details_file else None
    lock_file = Path(args.lock_file) if args.lock_file else None
    detail_log_dir = Path(args.detail_log_dir) if args.detail_log_dir else None
    if lock_file is None and details_file is not None:
        lock_file = details_file.with_name(f".{details_file.stem}.lock")

    try:
        if not rubrics_file.exists():
            write_outputs(
                result_file,
                build_reward(build_scorecard([])),
                details_file,
                {"error": "rubrics not found", "rubrics_file": str(rubrics_file)},
            )
            return

        rubrics_data = json.loads(rubrics_file.read_text(encoding="utf-8"))
        rubrics = rubrics_data.get("rubrics") or rubrics_data.get("evaluation_rubrics") or []
        if not rubrics:
            write_outputs(
                result_file,
                build_reward(build_scorecard([])),
                details_file,
                {"error": "no rubrics", "rubrics_file": str(rubrics_file)},
            )
            return

        total_rubric_count = len(rubrics)
        evaluated_model = args.evaluated_model or output_dir.name
        has_output_files = output_dir.exists() and any(path.is_file() for path in output_dir.rglob("*"))
        if not has_output_files:
            results = [
                build_failed_rubric_result(idx, rubric, "No output files found in the model output directory.")
                for idx, rubric in enumerate(rubrics)
            ]
        else:
            file_contents = extract_all_file_contents(output_dir)
            if not file_contents.strip():
                results = [
                    build_failed_rubric_result(idx, rubric, "Output files were unreadable or empty after conversion.")
                    for idx, rubric in enumerate(rubrics)
                ]
            else:
                existing_results = load_existing_rubric_results(details_file)
                api_base, api_key = resolve_api_config(args.judge_model, args.api_base, args.api_key)
                image_attachments = collect_image_attachments(output_dir)

                results: list[dict | None] = [existing_results.get(idx) for idx in range(len(rubrics))]
                with ThreadPoolExecutor(max_workers=args.max_workers) as executor:
                    futures = {
                        executor.submit(
                            judge_rubric,
                            idx,
                            rubric,
                            file_contents,
                            args.judge_model,
                            api_base,
                            api_key,
                            args.timeout_per_rubric,
                            args.max_retries,
                            image_attachments,
                        ): idx
                        for idx, rubric in enumerate(rubrics)
                        if results[idx] is None
                    }
                    for future in as_completed(futures):
                        idx = futures[future]
                        try:
                            results[idx], debug = future.result()
                        except Exception as exc:
                            results[idx] = build_failed_rubric_result(
                                idx,
                                rubrics[idx],
                                f"Judge raised an exception: {exc}",
                            )
                            debug = {
                                "api_base": api_base,
                                "parse_status": "failed",
                                "api_exit_code": 2,
                                "criterion_count": len(normalize_criteria(rubrics[idx])),
                                "criteria_list_text": "\n".join(
                                    f"Criterion {criterion_idx}: {criterion}"
                                    for criterion_idx, criterion in enumerate(normalize_criteria(rubrics[idx]))
                                ),
                                "rubric_text": rubrics[idx].get("rubric", ""),
                                "raw_response": "",
                                "error": str(exc),
                                "vision_used": False,
                                "attached_images": [],
                            }
                        if details_file is not None and lock_file is not None and results[idx] is not None:
                            append_rubric_result(
                                details_file,
                                lock_file,
                                evaluated_model,
                                args.judge_model,
                                total_rubric_count,
                                results[idx],
                            )
                        write_detail_log(
                            detail_log_dir,
                            args.detail_log_prefix,
                            idx,
                            args.judge_model,
                            debug,
                            results[idx],
                        )

                results = [result for result in results if result is not None]

        scorecard = build_scorecard(results)
        reward = build_reward(scorecard)
        details = build_details_report(
            evaluated_model,
            args.judge_model,
            results,
            total_count=total_rubric_count,
        )
        write_outputs(result_file, reward, details_file, details)
    except Exception as exc:
        fallback_reward = build_reward(build_scorecard([]))
        write_outputs(
            result_file,
            fallback_reward,
            details_file,
            {
                "error": str(exc),
                "traceback": traceback.format_exc(),
                "evaluated_model": args.evaluated_model or output_dir.name,
                "judge_model": args.judge_model,
            },
        )


if __name__ == "__main__":
    main()
