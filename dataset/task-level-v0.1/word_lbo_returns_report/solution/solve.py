#!/usr/bin/env python3
"""Golden solver: reads LBO DCF input.xlsx and generates LBO investment return analysis report."""
import os
import openpyxl
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

INPUT_PATH = '/workspace/data/input.xlsx'
OUTPUT_PATH = '/workspace/report.docx'

def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), '000000')
        borders.append(element)
    tblPr.append(borders)
    if tbl.tblPr is None:
        tbl.append(tblPr)

def add_page_number(doc):
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    run._r.append(fldChar1)
    run2 = para.add_run()
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' PAGE '
    run2._r.append(instrText)
    run3 = para.add_run()
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run3._r.append(fldChar2)

def format_paragraph(para):
    pf = para.paragraph_format
    pf.line_spacing = 1.5
    pf.first_line_indent = Pt(24)

def add_heading_formatted(doc, text, level):
    para = doc.add_paragraph()
    run = para.add_run(text)
    if level == 1:
        run.font.size = Pt(16)
        run.bold = True
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 2:
        run.font.size = Pt(14)
        run.bold = True
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif level == 3:
        run.font.size = Pt(12)
        run.bold = True
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(6)
    return para

def add_body_paragraph(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = 'FangSong'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'FangSong')
    run.font.size = Pt(12)
    format_paragraph(para)
    return para

def make_table_bold_header(table):
    for cell in table.rows[0].cells:
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True

def fmt_num(v, decimals=2):
    if v is None:
        return '-'
    try:
        return f'{float(v):,.{decimals}f}'
    except (TypeError, ValueError):
        return str(v)

def fmt_pct(v):
    if v is None:
        return '-'
    try:
        return f'{float(v)*100:.2f}%'
    except (TypeError, ValueError):
        return str(v)

def main():
    wb = openpyxl.load_workbook(INPUT_PATH, data_only=True)
    ws_input = wb['Input']
    ws_irr = wb['IRR_Returns']

    # Extract Input sheet data
    stock_price = ws_input.cell(8, 3).value
    offer_premium = ws_input.cell(9, 3).value
    offer_price = ws_input.cell(10, 3).value
    shares = ws_input.cell(11, 3).value
    equity_purchase = ws_input.cell(12, 3).value
    existing_debt = ws_input.cell(14, 3).value
    assumed_debt = ws_input.cell(15, 3).value
    transaction_costs = ws_input.cell(16, 3).value
    cash_infusion = ws_input.cell(18, 3).value
    total_funds = ws_input.cell(20, 3).value

    # Debt assumptions
    maturity_year = ws_input.cell(9, 9).value
    coupon_rate = ws_input.cell(11, 9).value
    avg_interest = ws_input.cell(12, 9).value
    optional_prepay = ws_input.cell(13, 9).value
    coupon_rate_2 = ws_input.cell(16, 9).value
    revolver_limit = ws_input.cell(17, 9).value
    undrawn_fee = ws_input.cell(18, 9).value

    # IRR Returns data
    exit_year = ws_irr.cell(4, 2).value

    # Cash flow rows, periods in columns D-M (4-13)
    period_cols = list(range(4, 14))
    display_cols = period_cols[:6]

    initial_inv = [ws_irr.cell(25, c).value for c in display_cols]
    cash_interest = [ws_irr.cell(26, c).value for c in display_cols]
    repayment = [ws_irr.cell(27, c).value for c in display_cols]
    principal = [ws_irr.cell(28, c).value for c in display_cols]
    participation = [ws_irr.cell(29, c).value for c in display_cols]
    irr_val = ws_irr.cell(30, 4).value

    wb.close()

    # Compute offer price if not available
    if offer_price is None and stock_price and offer_premium:
        offer_price = stock_price * (1 + offer_premium)
    if equity_purchase is None and offer_price and shares:
        equity_purchase = offer_price * shares

    # Create document
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)

    style = doc.styles['Normal']
    style.font.name = 'FangSong'
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'FangSong')

    add_page_number(doc)

    # Title
    add_heading_formatted(doc, 'LBO Investment Return Analysis Report', 1)

    # Section 1
    add_heading_formatted(doc, '1. Deal Overview and Return Summary', 2)
    add_body_paragraph(doc,
        'This report provides a detailed analysis of the investment returns from a Leveraged Buyout (LBO) '
        'transaction. An LBO is a private equity investment strategy that acquires a target company using '
        'significant debt financing, repays the debt from the company\'s future cash flows, and ultimately '
        'achieves equity investment returns through exit.')
    add_body_paragraph(doc,
        f'The target company\'s current stock price is ${fmt_num(stock_price, 2)}, '
        f'with an offer premium of {fmt_pct(offer_premium)}, '
        f'resulting in an offer price per share of ${fmt_num(offer_price, 2)}. '
        f'Shares outstanding total {fmt_num(shares, 0)}, '
        f'with a projected exit year of {int(exit_year) if exit_year else "N/A"}.')
    if irr_val:
        add_body_paragraph(doc,
            f'Based on current model assumptions, the project\'s Internal Rate of Return (IRR) is {fmt_pct(irr_val)}.')

    # Section 2
    add_heading_formatted(doc, '2. Deal Structure and Financing Arrangement', 2)
    add_body_paragraph(doc, 'The following table presents the basic structure and funding arrangement of this LBO transaction:')

    deal_data = [
        ['Item', 'Amount/Value'],
        ['Current Stock Price', f'${fmt_num(stock_price, 2)}'],
        ['Offer Premium', fmt_pct(offer_premium)],
        ['Offer Price Per Share', f'${fmt_num(offer_price, 2)}'],
        ['Shares Outstanding', fmt_num(shares, 0)],
        ['Equity Purchase Price', f'${fmt_num(equity_purchase, 2)}'],
        ['Assumed Debt', f'${fmt_num(assumed_debt, 2)}'],
        ['Transaction Costs', f'${fmt_num(transaction_costs, 2)}'],
        ['Cash Infusion', f'${fmt_num(cash_infusion, 2)}'],
        ['Total Funding Required', f'${fmt_num(total_funds, 2)}'],
    ]
    table = doc.add_table(rows=len(deal_data), cols=2)
    table.style = 'Table Grid'
    set_table_borders(table)
    for i, row_data in enumerate(deal_data):
        for j, val in enumerate(row_data):
            table.rows[i].cells[j].text = val
    make_table_bold_header(table)

    # Section 3
    add_heading_formatted(doc, '3. Debt Terms', 2)
    add_body_paragraph(doc, 'The debt financing arrangement for this transaction is as follows:')

    debt_data = [
        ['Item', 'Value'],
        ['Maturity', f'{int(maturity_year) if maturity_year else "-"} years'],
        ['Coupon Rate (Senior Debt)', fmt_pct(coupon_rate)],
        ['Coupon Rate (Subordinated Debt)', fmt_pct(coupon_rate_2)],
        ['Revolving Credit Facility', f'${fmt_num(revolver_limit, 2)}'],
        ['Undrawn Commitment Fee', fmt_pct(undrawn_fee)],
        ['Optional Prepayment', 'Yes' if optional_prepay else 'No'],
    ]
    table = doc.add_table(rows=len(debt_data), cols=2)
    table.style = 'Table Grid'
    set_table_borders(table)
    for i, row_data in enumerate(debt_data):
        for j, val in enumerate(row_data):
            table.rows[i].cells[j].text = val
    make_table_bold_header(table)

    # Section 4
    add_heading_formatted(doc, '4. Cash Flow and IRR Analysis', 2)
    add_body_paragraph(doc,
        'The following table shows the cash flow allocation for each post-acquisition period, '
        'including initial investment, interest income, debt repayment, principal recovery, '
        'and exit participation proceeds:')

    cf_headers = ['Item'] + [f'Period {i+1}' for i in range(len(display_cols))]
    cf_data = [
        cf_headers,
        ['Initial Investment'] + [fmt_num(v) for v in initial_inv],
        ['Cash Interest'] + [fmt_num(v) for v in cash_interest],
        ['Repayment'] + [fmt_num(v) for v in repayment],
        ['Principal'] + [fmt_num(v) for v in principal],
        ['Participation on Exit'] + [fmt_num(v) for v in participation],
    ]
    table = doc.add_table(rows=len(cf_data), cols=len(cf_headers))
    table.style = 'Table Grid'
    set_table_borders(table)
    for i, row_data in enumerate(cf_data):
        for j, val in enumerate(row_data):
            if j < len(table.rows[i].cells):
                table.rows[i].cells[j].text = val
    make_table_bold_header(table)

    # Cash flow chart
    ci_vals = [v for v in cash_interest if v is not None and isinstance(v, (int, float))]
    rep_vals = [v for v in repayment if v is not None and isinstance(v, (int, float))]
    if ci_vals or rep_vals:
        fig, ax = plt.subplots(figsize=(6, 3.5))
        x = range(1, max(len(ci_vals), len(rep_vals)) + 1)
        if ci_vals:
            ax.bar([i - 0.2 for i in x[:len(ci_vals)]], ci_vals, width=0.35,
                   label='Cash Interest', color='#4472C4')
        if rep_vals:
            ax.bar([i + 0.2 for i in x[:len(rep_vals)]], rep_vals, width=0.35,
                   label='Repayment', color='#ED7D31')
        ax.set_xlabel('Period', fontsize=10)
        ax.set_ylabel('Amount ($)', fontsize=10)
        ax.set_title('Cash Flow Components by Period', fontsize=12)
        ax.legend()
        ax.grid(axis='y', alpha=0.3)
        plt.tight_layout()
        chart_path = '/tmp/lbo_cashflow.png'
        plt.savefig(chart_path, dpi=150)
        plt.close()
        doc.add_picture(chart_path, width=Inches(5.0))

    # Section 5
    add_heading_formatted(doc, '5. Exit Strategy and Sensitivity Analysis', 2)
    add_body_paragraph(doc,
        f'The model assumes an exit in {int(exit_year) if exit_year else "N/A"}. '
        'Potential exit methods include IPO (Initial Public Offering), strategic sale, or secondary LBO. '
        'The enterprise value at exit is estimated based on the exit-year EBITDA multiple.')
    add_body_paragraph(doc, 'Key factors affecting IRR include:')

    factors = [
        'Exit timing: Delayed exit reduces IRR as the holding period extends while total return growth is limited',
        'Exit EBITDA multiple: Each 1x increase in the multiple raises IRR by approximately 3-5 percentage points',
        'Debt repayment pace: Accelerated repayment increases equity proportion at exit but reduces tax shield benefits',
        'Operating performance: EBITDA growth is the core driver of LBO returns',
        'Interest rate risk: Higher floating-rate debt exposure means rate hikes compress distributable cash flow',
        'Refinancing risk: The ability to refinance maturing debt on reasonable terms affects project sustainability',
    ]
    for factor in factors:
        para = doc.add_paragraph()
        para.style = 'List Bullet'
        run = para.add_run(factor)
        run.font.name = 'FangSong'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'FangSong')
        run.font.size = Pt(12)

    doc.save(OUTPUT_PATH)
    print(f'Report saved to {OUTPUT_PATH}')

# [oracle] 把本脚本源码复制到 /workspace/generated_code.py,供 L2 code judge 读取
# (判分器 _find_code_file 优先读该路径;solve.py 本身不在 /workspace,不复制则 L2 找不到代码 → L2=0)
def _emit_generated_code():
    import shutil, os
    try:
        os.makedirs('/workspace', exist_ok=True)
        shutil.copyfile(os.path.abspath(__file__), '/workspace/generated_code.py')
    except Exception as e:
        print(f'warn: could not emit generated_code.py: {e}')


if __name__ == '__main__':
    _emit_generated_code()
    main()
